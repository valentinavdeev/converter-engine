"""Local image embedding and hidden-slide SVG capture contracts."""

from __future__ import annotations

import base64
import io
from pathlib import Path
from urllib.parse import quote

import pytest
from PIL import Image
from pptx import Presentation

from html_to_pptx import convert
from html_to_pptx.converter import (
    EXTRACTION_JS, PIXELS_TO_INCHES_X, PIXELS_TO_INCHES_Y,
    _prepare_slide_images, _rasterize_inline_svgs, _walk_elements,
)


@pytest.mark.asyncio
async def test_local_image_survives_full_conversion(tmp_path: Path):
    image_path = tmp_path / "synthetic image.png"
    Image.new("RGB", (32, 24), (25, 100, 200)).save(image_path)
    source = tmp_path / "deck # local.html"
    source.write_text(
        '<style>.slide {width:1280px;height:720px}</style>'
        '<section class="slide"><img src="synthetic image.png" '
        'style="width:64px;height:48px"></section>',
        encoding="utf-8",
    )
    output = tmp_path / "local.pptx"
    await convert(str(source), str(output))
    presentation = Presentation(str(output))
    pictures = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "image")]
    assert len(pictures) == 1
    assert pictures[0].image.blob == image_path.read_bytes()
    assert pictures[0].image.size == (32, 24)



@pytest.mark.asyncio
@pytest.mark.parametrize("source_kind", ["local", "base64", "percent_encoded"])
async def test_svg_image_becomes_positioned_png_picture(tmp_path: Path, source_kind: str):
    svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20">'
        '<rect width="10" height="20" fill="#19c864"/>'
        '<rect x="10" width="10" height="20" fill="#1964c8"/>'
        '</svg>'
    )
    if source_kind == "local":
        (tmp_path / "icon.svg").write_text(svg, encoding="utf-8")
        image_source = "icon.svg"
    elif source_kind == "base64":
        image_source = "data:image/svg+xml;base64," + base64.b64encode(svg.encode()).decode()
    else:
        image_source = "data:image/svg+xml," + quote(svg)
    source = tmp_path / "svg-image.html"
    source.write_text(
        '<style>body {margin:0}.slide {width:1920px;height:1080px;position:relative}</style>'
        f'<section class="slide" style="display:none"><img src="{image_source}" width="100" height="100" '
        'style="position:absolute;left:240px;top:160px"></section>',
        encoding="utf-8",
    )
    output = tmp_path / "svg-image.pptx"
    await convert(str(source), str(output))
    pictures = [
        shape for shape in Presentation(str(output)).slides[0].shapes
        if hasattr(shape, "image")
    ]
    assert len(pictures) == 1
    picture = pictures[0]
    assert picture.left / 914400 == pytest.approx(240 * PIXELS_TO_INCHES_X)
    assert picture.top / 914400 == pytest.approx(160 * PIXELS_TO_INCHES_Y)
    assert picture.width / 914400 == pytest.approx(100 * PIXELS_TO_INCHES_X)
    assert picture.height / 914400 == pytest.approx(100 * PIXELS_TO_INCHES_Y)
    with Image.open(io.BytesIO(picture.image.blob)) as image:
        assert image.format == "PNG"
        assert image.size == (20, 20)
        assert image.convert("RGB").getpixel((5, 10)) == (25, 200, 100)
        assert image.convert("RGB").getpixel((15, 10)) == (25, 100, 200)


@pytest.mark.asyncio
async def test_nonparticipating_images_do_not_abort_export(tmp_path: Path):
    image_path = tmp_path / "visible.png"
    Image.new("RGB", (32, 24), (25, 100, 200)).save(image_path)
    source = tmp_path / "ignored-images.html"
    source.write_text(
        '<style>body {margin:0}.slide {width:1920px;height:1080px;position:relative}</style>'
        '<img><img src="outside-missing.png">'
        '<section class="slide">'
        '<img src="visible.png" width="64" height="48">'
        '<img style="display:none"><img src="hidden-missing.png" style="display:none">'
        '<div style="display:none"><img src="ancestor-hidden.png"></div>'
        '<div style="visibility:hidden"><img src="ancestor-invisible.png" '
        'style="visibility:visible"></div>'
        '<div style="opacity:0"><img src="ancestor-transparent.png"></div>'
        '<img src="off-slide-missing.png" width="32" height="24" '
        'style="position:absolute;left:2200px">'
        '<div style="position:absolute;left:2200px;width:100px;height:100px">'
        '<img src="pruned-wrapper.png" width="32" height="24" '
        'style="position:absolute;left:-2200px"></div>'
        '</section>'
        '<div style="display:none"><section class="slide">'
        '<img src="outside-hidden-ancestor.png"></section></div>',
        encoding="utf-8",
    )
    output = tmp_path / "ignored-images.pptx"
    await convert(str(source), str(output))
    presentation = Presentation(str(output))
    assert len(presentation.slides) == 2
    pictures = [
        shape for shape in presentation.slides[0].shapes if hasattr(shape, "image")
    ]
    assert len(pictures) == 1
    assert pictures[0].image.blob == image_path.read_bytes()
    assert not any(hasattr(shape, "image") for shape in presentation.slides[1].shapes)


@pytest.mark.asyncio
async def test_initially_inactive_slides_prepare_selected_images(tmp_path: Path):
    first = tmp_path / "first.png"
    selected = tmp_path / "selected.png"
    Image.new("RGB", (32, 24), (25, 100, 200)).save(first)
    Image.new("RGB", (32, 24), (200, 100, 25)).save(selected)
    embedded = "data:image/png;base64," + base64.b64encode(selected.read_bytes()).decode()
    source = tmp_path / "inactive-images.html"
    source.write_text(
        '<style>.slide {width:1920px;height:1080px;display:none}'
        '.slide.active {display:block}</style>'
        '<section class="slide active"><img src="first.png" width="64" height="48"></section>'
        '<section class="slide"><picture><source srcset="selected.png">'
        '<img src="missing-fallback.png" width="64" height="48"></picture></section>'
        '<section class="slide" style="display:none"><picture>'
        f'<source srcset="{embedded}"><img src="missing-embedded-fallback.png" '
        'width="64" height="48"></picture></section>',
        encoding="utf-8",
    )
    output = tmp_path / "inactive-images.pptx"
    await convert(str(source), str(output))
    presentation = Presentation(str(output))
    assert len(presentation.slides) == 3
    for slide, expected in zip(presentation.slides, [first, selected, selected]):
        pictures = [shape for shape in slide.shapes if hasattr(shape, "image")]
        assert len(pictures) == 1
        assert pictures[0].image.blob == expected.read_bytes()


@pytest.mark.asyncio
@pytest.mark.parametrize("missing", [False, True])
async def test_image_preparation_restores_slide_state(tmp_path: Path, missing: bool):
    from playwright.async_api import async_playwright

    if not missing:
        Image.new("RGB", (32, 24), (25, 100, 200)).save(tmp_path / "picture.png")
    source = tmp_path / "state.html"
    source.write_text(
        '<style>.slide {width:1920px;height:1080px;display:none}'
        '.slide.active {display:grid}</style>'
        '<section class="slide active authored" style="color:red"></section>'
        '<section class="slide authored" style="display:none;color:blue">'
        '<img src="picture.png" width="64" height="48"></section>',
        encoding="utf-8",
    )
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(headless=True)
        try:
            page = await browser.new_page()
            await page.goto(source.as_uri(), wait_until="networkidle")
            snapshot = (
                "() => Array.from(document.querySelectorAll('.slide'), "
                "slide => [slide.getAttribute('class'), slide.getAttribute('style')])"
            )
            original = await page.evaluate(snapshot)
            if missing:
                with pytest.raises(FileNotFoundError):
                    await _prepare_slide_images(page)
            else:
                await _prepare_slide_images(page)
                measurements = await page.evaluate(EXTRACTION_JS)
                pictures = [
                    element for element in _walk_elements(measurements[1]["elements"])
                    if element.get("isImage")
                ]
                assert len(pictures) == 1
                assert pictures[0]["src"].startswith("data:image/png;base64,")
            assert await page.evaluate(snapshot) == original
        finally:
            await browser.close()

@pytest.mark.asyncio
async def test_remote_image_is_rejected_without_output(tmp_path: Path):
    source = tmp_path / "remote.html"
    source.write_text(
        '<section class="slide"><img src="https://example.invalid/image.png"></section>',
        encoding="utf-8",
    )
    output = tmp_path / "remote.pptx"
    with pytest.raises(ValueError, match="Unsupported image source"):
        await convert(str(source), str(output))
    assert not output.exists()


@pytest.mark.asyncio
async def test_svg_capture_preserves_authored_layout_and_slide_state(tmp_path: Path):
    from playwright.async_api import async_playwright

    html = """
        <style>
            .slide { width:1280px; height:720px; gap:20px; }
            .grid { display:grid; grid-template-columns:200px 80px; align-items:start; }
            .flex { display:flex; align-items:flex-start; }
            .slide > div { width:100px; height:40px; flex-shrink:0; }
            .slide svg { width:100%; height:40px; }
            .flex svg { width:80px; flex-shrink:0; }
        </style>
        <section class="slide grid preserved" style="display:none;color:red">
            <div></div><svg viewBox="0 0 80 40"><rect width="80" height="40" fill="#19c864"/></svg>
        </section>
        <section class="slide flex preserved" style="display:none;color:blue">
            <div></div><svg viewBox="0 0 80 40"><rect width="80" height="40" fill="#19c864"/></svg>
        </section>
    """
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(headless=True)
        try:
            page = await browser.new_page(viewport={"width": 1280, "height": 720})
            await page.set_content(html)
            snapshot_js = "() => Array.from(document.querySelectorAll('.slide'), s => [s.getAttribute('style'), s.getAttribute('class')])"
            original = await page.evaluate(snapshot_js)
            measurements = await page.evaluate(EXTRACTION_JS)
            assert await page.evaluate(snapshot_js) == original
            svgs = [
                next(element for element in _walk_elements(slide["elements"]) if element.get("isSvg"))
                for slide in measurements
            ]
            assert [svg["x"] for svg in svgs] == [220, 120]
            assert [svg["width"] for svg in svgs] == [80, 80]
            await _rasterize_inline_svgs(page, measurements)
            assert await page.evaluate(snapshot_js) == original
            for svg in svgs:
                assert svg["isImage"] is True
                assert svg["isSvg"] is False
                pixels = base64.b64decode(svg["src"].split(",", 1)[1])
                with Image.open(io.BytesIO(pixels)) as image:
                    assert image.size == (80, 40)
                    assert image.convert("RGB").getpixel((40, 20)) == (25, 200, 100)
        finally:
            await browser.close()
    source = tmp_path / "inline-svg.html"
    source.write_text(html, encoding="utf-8")
    output = tmp_path / "inline-svg.pptx"
    await convert(str(source), str(output))
    presentation = Presentation(str(output))
    assert len(presentation.slides) == 2
    for slide in presentation.slides:
        pictures = [shape for shape in slide.shapes if hasattr(shape, "image")]
        assert len(pictures) == 1
        assert pictures[0].image.size == (80, 40)


@pytest.mark.asyncio
@pytest.mark.parametrize("alias", ["same", "symlink", "hardlink"])
async def test_convert_rejects_source_as_output(tmp_path: Path, alias: str):
    source = tmp_path / "source.html"
    original = b'<section class="slide">Preserved</section>'
    source.write_bytes(original)
    output = source
    if alias != "same":
        output = tmp_path / "alias.pptx"
        if alias == "symlink":
            output.symlink_to(source)
        else:
            output.hardlink_to(source)
    with pytest.raises(ValueError, match="Input and output must be different files"):
        await convert(str(source), str(output))
    assert source.read_bytes() == original
    assert output.read_bytes() == original
