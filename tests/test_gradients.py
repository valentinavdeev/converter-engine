"""Gradient fidelity checked through saved and reopened native PPTX objects."""

from __future__ import annotations

from io import BytesIO

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from html_to_pptx import converter as C


def _saved_gradient(css: str, target: str = "shape", opacity: float = 1.0):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if target == "text":
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "Synthetic gradient text"
        assert C._apply_gradient_text(run, css, opacity)
    elif target == "slide":
        assert C._apply_gradient_fill(slide.background._element, css, opacity=opacity)
    else:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1)
        )
        assert C._apply_gradient_fill(shape._element, css, opacity=opacity)
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    restored = Presentation(stream).slides[0]
    if target == "text":
        ancestor = restored.shapes[0].text_frame.paragraphs[0].runs[0]._r
    elif target == "slide":
        ancestor = restored.background._element
    else:
        ancestor = restored.shapes[0]._element
    return ancestor.find(".//" + qn("a:gradFill"))


def _stops(gradient):
    result = []
    for stop in gradient.find(qn("a:gsLst")):
        color = stop.find(qn("a:srgbClr"))
        alpha = color.find(qn("a:alpha"))
        result.append((int(stop.get("pos")), color.get("val"),
                       int(alpha.get("val")) if alpha is not None else 100000))
    return result


@pytest.mark.parametrize("target", ["shape", "slide", "text"])
def test_off_center_stop_retains_color_alpha_and_angle_after_reopen(target):
    gradient = _saved_gradient(
        "linear-gradient(90deg, rgba(255, 0, 0, 1) 0%, "
        "rgba(0, 255, 0, 0.25) 30%, rgba(0, 0, 255, 1) 100%)",
        target, opacity=0.5,
    )
    assert _stops(gradient) == [
        (0, "FF0000", 50000), (30000, "00FF00", 12500), (100000, "0000FF", 50000),
    ]
    assert gradient.find(qn("a:lin")).get("ang") == "0"


@pytest.mark.parametrize(
    ("css", "positions"),
    [
        (
            "linear-gradient(rgb(0,0,0), rgb(50,50,50), rgb(100,100,100) 20%, "
            "rgb(150,150,150), rgb(255,255,255))",
            [0, 10000, 20000, 60000, 100000],
        ),
        (
            "linear-gradient(rgb(0,0,0) 10%, rgb(50,50,50), rgb(100,100,100), "
            "rgb(150,150,150) 70%, rgb(255,255,255) 100%)",
            [10000, 30000, 50000, 70000, 100000],
        ),
        (
            "linear-gradient(rgb(0,0,0) 60%, rgb(50,50,50), rgb(100,100,100) 20%, "
            "rgb(150,150,150) 40%, rgb(255,255,255))",
            [60000, 60000, 60000, 60000, 100000],
        ),
    ],
)
def test_omitted_positions_and_css_monotonic_fixup_survive_reopen(css, positions):
    stops = _stops(_saved_gradient(css))
    assert [stop[0] for stop in stops] == positions
    assert [stop[1] for stop in stops] == ["000000", "323232", "646464", "969696", "FFFFFF"]


@pytest.mark.parametrize("target", ["shape", "text"])
def test_repeated_positions_keep_both_sides_of_hard_transition(target):
    gradient = _saved_gradient(
        "linear-gradient(to bottom, rgb(255,0,0) 0% 40%, rgb(0,0,255) 40% 100%)",
        target,
    )
    assert _stops(gradient) == [
        (0, "FF0000", 100000), (40000, "FF0000", 100000),
        (40000, "0000FF", 100000), (100000, "0000FF", 100000),
    ]
    assert gradient.find(qn("a:lin")).get("ang") == "5400000"


def test_explicit_boundary_transitions_are_not_coalesced():
    gradient = _saved_gradient(
        "linear-gradient(rgb(255,0,0) 0%, rgb(0,255,0) 0%, "
        "rgb(0,255,0) 100%, rgb(0,0,255) 100%)"
    )
    assert _stops(gradient) == [
        (0, "FF0000", 100000), (0, "00FF00", 100000),
        (100000, "00FF00", 100000), (100000, "0000FF", 100000),
    ]


def test_outside_percentage_positions_clip_without_rescaling_colors():
    gradient = _saved_gradient(
        "linear-gradient(rgb(255,0,0) -50%, rgb(0,0,255) 150%)"
    )
    assert _stops(gradient) == [(0, "BF0040", 100000), (100000, "4000BF", 100000)]


def test_clipped_transparency_uses_premultiplied_color():
    gradient = _saved_gradient(
        "linear-gradient(transparent -100%, rgba(255,0,0,1) 100%)"
    )
    assert _stops(gradient) == [(0, "FF0000", 50000), (100000, "FF0000", 100000)]


@pytest.mark.parametrize(
    ("positions", "color"),
    [("-100%, rgb(0,0,255) -50%", "0000FF"),
     ("150%, rgb(0,0,255) 200%", "FF0000")],
)
def test_stops_entirely_outside_visible_range_become_correct_constant_fill(positions, color):
    gradient = _saved_gradient(f"linear-gradient(rgb(255,0,0) {positions})")
    assert _stops(gradient) == [(0, color, 100000), (100000, color, 100000)]


def test_modern_rgb_percentage_channels_and_signed_angle_preserve_native_alpha():
    gradient = _saved_gradient(
        "linear-gradient(-90deg, rgb(100% 0% 0% / 50%) 0, rgba(0 0 255 / 0.25) 100%)",
        "text",
    )
    assert _stops(gradient) == [(0, "FF0000", 50000), (100000, "0000FF", 25000)]
    assert gradient.find(qn("a:lin")).get("ang") == "10800000"


@pytest.mark.parametrize(
    "css",
    [
        "linear-gradient(rgb(0,0,0) 10px, rgb(255,255,255) 90px)",
        "linear-gradient(rgb(0,0,0) calc(10% + 5px), rgb(255,255,255))",
        "linear-gradient(rgb(0,0,0), 30%, rgb(255,255,255))",
        "linear-gradient(rgb(0,0,0), rgb(255,255,255)), url(example.png)",
        "radial-gradient(rgb(0,0,0), rgb(255,255,255))",
        "conic-gradient(rgb(0,0,0), rgb(255,255,255))",
        "linear-gradient(rgb(0,0,0), rgb(255,255,255)), linear-gradient(rgb(1,1,1), rgb(2,2,2))",
    ],
)
def test_unsupported_syntax_does_not_overwrite_existing_fill(css):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(12, 34, 56)
    assert not C._apply_gradient_fill(shape._element, css)
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    restored = Presentation(stream).slides[0].shapes[0]
    assert restored._element.find(".//" + qn("a:gradFill")) is None
    assert restored.fill.fore_color.rgb == RGBColor(12, 34, 56)
