"""Consumer-visible native-object regressions; browser/layout smoke uses fixture.html."""
import base64
from copy import deepcopy
from io import BytesIO
import unittest

from html_to_pptx import render_pptx
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def element(**values):
    result = dict(tag='div', x=80, y=100, width=200, height=100,
                  text='', children=[], backgroundColor='rgba(0, 0, 0, 0)',
                  backgroundImage='none', opacity=1, rotation=0,
                  fontFamily='Arial', fontSize=50, fontWeight='400',
                  fontStyle='normal', lineHeight='50px', whiteSpace='pre',
                  letterSpacing=0, color='rgb(0, 0, 0)', textAlign='left',
                  paddingTop=0, paddingBottom=0, paddingLeft=0, paddingRight=0)
    result.update(values)
    return result


def export(*elements):
    deck = render_pptx([dict(backgroundColor='rgb(229, 237, 244)', elements=list(elements))])
    data = BytesIO()
    deck.save(data)
    data.seek(0)
    return Presentation(data).slides[0]


def leaves(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from leaves(shape.shapes)
        else:
            yield shape


class NativeFidelityTests(unittest.TestCase):
    def test_half_pixel_rule_remains_visible_native_geometry(self):
        slide = export(element(width=120, height=.5, backgroundColor='rgb(53, 69, 82)'))
        self.assertEqual(len(slide.shapes), 1)
        shape = slide.shapes[0]
        self.assertEqual(shape.shape_type, MSO_SHAPE_TYPE.AUTO_SHAPE)
        self.assertAlmostEqual(shape.height / 6350, .5, places=3)
        self.assertEqual(str(shape.fill.fore_color.rgb), '354552')

    def test_text_size_and_tracking_survive_tight_browser_box(self):
        slide = export(element(tag='span', text='22,7', width=91.172, height=50, letterSpacing=-1.54))
        shape = next(s for s in slide.shapes if s.has_text_frame and s.text == '22,7')
        run = shape.text_frame.paragraphs[0].runs[0]
        self.assertAlmostEqual(run.font.size.pt, 25, places=2)
        self.assertEqual(run._r.rPr.get('spc'), '-77')
        self.assertFalse(shape.text_frame.word_wrap)

    def test_preformatted_boundary_spaces_and_newlines_survive_save(self):
        original = ' A  B \n C '
        slide = export(element(tag='span', text=original, width=240, height=100, fontSize=30, lineHeight='30px'))
        texts = [s.text.replace('\v', '\n') for s in slide.shapes if s.has_text_frame and s.text]
        self.assertEqual(texts, [original])

    def test_fractional_small_font_is_not_clamped_or_shrunk(self):
        slide = export(element(tag='span', text='8,3', width=18.078, height=13, fontSize=13, lineHeight='13px', fontWeight='700'))
        shape = next(s for s in slide.shapes if s.has_text_frame and s.text == '8,3')
        self.assertEqual(shape.text_frame.paragraphs[0].runs[0].font.size.pt, 6.5)

    def test_transparent_overlay_keeps_image_and_fill_separately_editable(self):
        png = BytesIO()
        Image.new('RGB', (2, 2), (19, 51, 87)).save(png, format='PNG')
        src = 'data:image/png;base64,' + base64.b64encode(png.getvalue()).decode()
        picture = element(tag='img', isImage=True, src=src, width=400, height=300)
        overlay = element(width=400, height=300, backgroundColor='rgba(255, 255, 255, 0.3)')
        before = deepcopy([picture, overlay])
        slide = export(picture, overlay)
        self.assertEqual([picture, overlay], before)
        image = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
        self.assertEqual(image.image.blob, png.getvalue())
        shape = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
        self.assertEqual(str(shape.fill.fore_color.rgb), 'FFFFFF')
        self.assertEqual(shape._element.xpath('./p:spPr/a:solidFill/a:srgbClr/a:alpha')[0].get('val'), '30000')

    def test_theme_does_not_add_unauthored_shadow(self):
        slide = export(element(backgroundColor='rgb(255, 255, 255)', borderRadius='20px'))
        shape = slide.shapes[0]
        refs = shape._element.xpath('./p:style/a:effectRef')
        self.assertTrue(all(r.get('idx') == '0' for r in refs))
        self.assertFalse(shape._element.xpath('.//a:outerShdw'))

    def test_explicit_nested_groups_preserve_geometry_and_paint_order(self):
        inner = element(isGroup=True, groupName='Inner', children=[
            element(x=200, y=220, width=60, height=50, backgroundColor='rgb(1, 2, 3)'),
            element(tag='span', x=205, y=225, width=50, height=20, fontSize=20, lineHeight='20px', text='B')])
        outer = element(isGroup=True, groupName='Outer', children=[
            element(x=100, y=150, width=300, height=300, backgroundColor='rgb(255, 255, 255)'), inner])
        slide = export(outer, element(x=500, backgroundColor='rgb(255, 0, 0)'))
        self.assertEqual(slide.shapes[0].shape_type, MSO_SHAPE_TYPE.GROUP)
        group = slide.shapes[0]
        self.assertEqual(group.name, 'Outer')
        self.assertEqual(group.shapes[1].shape_type, MSO_SHAPE_TYPE.GROUP)
        self.assertEqual(group.shapes[1].name, 'Inner')
        nested_shape = group.shapes[1].shapes[0]
        self.assertAlmostEqual(nested_shape.left / 6350, 200, places=3)
        self.assertAlmostEqual(nested_shape.top / 6350, 220, places=3)
        self.assertEqual([s.text for s in leaves(group.shapes) if s.has_text_frame and s.text], ['B'])
        self.assertEqual(str(slide.shapes[1].fill.fore_color.rgb), 'FF0000')


if __name__ == '__main__':
    unittest.main()
