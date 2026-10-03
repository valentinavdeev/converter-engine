"""Saved-PPTX regressions for reviewed inline and measured-text boundaries."""
from io import BytesIO

import pytest
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from html_to_pptx import convert, extract_measurements, render_pptx


CSS_PX_TO_EMU = 6350


def html_file(tmp_path, content):
    source = tmp_path / "review.html"
    source.write_text(
        '<!doctype html><meta charset="utf-8"><style>'
        '*{box-sizing:border-box}body{margin:0}'
        '.slide{position:relative;width:1920px;height:1080px;'
        'background:white;font-family:Arial;display:block}'
        '</style><section class="slide">' + content + '</section>',
        encoding="utf-8",
    )
    return source


def reopen(measurements):
    stream = BytesIO()
    render_pptx(measurements).save(stream)
    stream.seek(0)
    return Presentation(stream).slides[0]


async def measured_slide(tmp_path, content):
    measurements = await extract_measurements(str(html_file(tmp_path, content)))
    return measurements[0]["elements"], reopen(measurements)


def text_shapes(slide):
    return [shape for shape in slide.shapes if shape.has_text_frame and shape.text]


@pytest.mark.asyncio
@pytest.mark.parametrize("border", ["border:2px solid red", "border-left:2px solid red"])
async def test_border_only_mixed_paragraph_keeps_native_border(tmp_path, border):
    elements, slide = await measured_slide(
        tmp_path,
        '<p style="margin:0;width:400px;height:80px;font-size:40px;' + border + '">'
        'Hello <strong>world</strong></p>',
    )
    assert [shape.text for shape in text_shapes(slide)] == ["Hello world"]
    runs = text_shapes(slide)[0].text_frame.paragraphs[0].runs
    assert [(run.text, run.font.bold) for run in runs] == [("Hello ", False), ("world", True)]
    artwork = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    if border.startswith("border-left"):
        accent = next(shape for shape in artwork if shape.fill.type == MSO_FILL_TYPE.SOLID
                      and str(shape.fill.fore_color.rgb) == "FF0000")
        assert accent.left / CSS_PX_TO_EMU == pytest.approx(elements[0]["x"], abs=.001)
        assert accent.height / CSS_PX_TO_EMU == pytest.approx(80, abs=.001)
    else:
        outline = next(shape for shape in artwork if str(shape.line.color.rgb) == "FF0000")
        assert outline.line.width.pt == pytest.approx(1)
        assert outline.width / CSS_PX_TO_EMU == pytest.approx(400, abs=.001)
        assert outline.height / CSS_PX_TO_EMU == pytest.approx(80, abs=.001)


@pytest.mark.asyncio
async def test_decorative_inline_block_survives_without_duplicate_text(tmp_path):
    elements, slide = await measured_slide(
        tmp_path,
        '<h4 style="margin:0;font-size:40px"><span style="display:inline-block;'
        'width:12px;height:12px;background:red"></span>Tokyo</h4>',
    )
    text = text_shapes(slide)
    assert [shape.text for shape in text] == ["Tokyo"]
    artwork = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
    assert str(artwork.fill.fore_color.rgb) == "FF0000"
    assert artwork.width / CSS_PX_TO_EMU == pytest.approx(12, abs=.001)
    assert artwork.height / CSS_PX_TO_EMU == pytest.approx(12, abs=.001)
    measured = elements[0]
    assert text[0].left / CSS_PX_TO_EMU == pytest.approx(
        measured["x"] + measured["textGeometry"]["x"], abs=.001,
    )
    assert text[0].left >= artwork.left + artwork.width


@pytest.mark.asyncio
async def test_nested_inline_background_and_border_remain_separate_from_text(tmp_path):
    _, slide = await measured_slide(
        tmp_path,
        '<p style="margin:0;font-size:40px">Hello <strong style="background:red;'
        'border:2px solid blue"><span style="background:yellow">world</span></strong>!</p>',
    )
    assert [shape.text for shape in text_shapes(slide)] == ["Hello world!"]
    artwork = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert [str(shape.fill.fore_color.rgb) for shape in artwork] == ["FF0000", "FFFF00"]
    assert str(artwork[0].line.color.rgb) == "0000FF"
    assert artwork[0].line.width.pt == pytest.approx(1)
    assert slide.shapes[-1].text == "Hello world!"


@pytest.mark.asyncio
@pytest.mark.parametrize("mixed", [False, True])
async def test_centered_preformatted_flex_multiline_uses_measured_placement(tmp_path, mixed):
    content = '<span>A\n<strong>B</strong></span>' if mixed else 'A\nB'
    elements, slide = await measured_slide(
        tmp_path,
        '<div style="display:flex;justify-content:center;align-items:center;'
        'width:600px;height:300px;font-size:40px;white-space:pre">' + content + '</div>',
    )
    measured = elements[0]
    geometry = measured["textGeometry"]
    assert geometry["lineCount"] == 2
    assert geometry["x"] > 250
    text = text_shapes(slide)[0]
    assert text.text.replace('\v', '\n') == 'A\nB'
    assert text.left / CSS_PX_TO_EMU == pytest.approx(measured["x"] + geometry["x"], abs=.001)
    assert text.top / CSS_PX_TO_EMU == pytest.approx(
        measured["y"] + geometry["baseline"] - geometry["ascent"], abs=.001,
    )
    assert text.width / CSS_PX_TO_EMU == pytest.approx(geometry["width"], abs=.001)
    assert text.text_frame.vertical_anchor == MSO_ANCHOR.TOP
    assert not text.text_frame.word_wrap


@pytest.mark.asyncio
@pytest.mark.parametrize("align, expected", [("center", PP_ALIGN.CENTER), ("right", PP_ALIGN.RIGHT)])
async def test_normal_multiline_keeps_measured_union_and_paragraph_alignment(tmp_path, align, expected):
    elements, slide = await measured_slide(
        tmp_path,
        '<p style="margin:0;width:600px;font-size:40px;padding:20px;text-align:'
        + align + '">A<br>Longer line</p>',
    )
    measured = elements[0]
    geometry = measured["textGeometry"]
    assert geometry["lineCount"] == 2
    text = text_shapes(slide)[0]
    assert text.text.replace('\v', '\n') == 'A\nLonger line'
    assert text.left / CSS_PX_TO_EMU == pytest.approx(measured["x"] + geometry["x"], abs=.001)
    assert text.width / CSS_PX_TO_EMU == pytest.approx(geometry["width"], abs=.001)
    assert text.text_frame.margin_left == text.text_frame.margin_right == 0
    assert text.text_frame.paragraphs[0].alignment == expected
    assert text.text_frame.vertical_anchor == MSO_ANCHOR.TOP


@pytest.mark.asyncio
@pytest.mark.parametrize("mixed", [False, True])
async def test_one_pixel_font_reports_actionable_html_context(tmp_path, mixed):
    content = ('<p style="font-size:40px">Normal <strong id="tiny" style="font-size:1px">'
               'small</strong></p>') if mixed else '<span id="tiny" style="font-size:1px">small</span>'
    source = html_file(tmp_path, content)
    target = tmp_path / "rejected.pptx"
    original = source.read_bytes()
    with pytest.raises(ValueError) as failure:
        await convert(str(source), str(target))
    message = str(failure.value)
    assert ("strong#tiny" if mixed else "span#tiny") in message
    assert "small" in message
    assert not target.exists()
    assert source.read_bytes() == original


def sized_element(size_px, mixed):
    element = dict(tag="p", sourceSelector="p#bounds", x=100, y=100, width=500,
                   height=100, text="Boundary", children=[], fontSize=size_px,
                   color="rgb(0,0,0)", fontFamily="Arial", lineHeight="normal")
    if mixed:
        element["text"] = ""
        element["inlineRuns"] = [dict(text="Boundary", fontSize=size_px,
                                      sourceSelector="span#bounds")]
    return element


@pytest.mark.parametrize("mixed", [False, True])
@pytest.mark.parametrize("size_px", [1.99, 8000.02])
def test_font_sizes_outside_native_bounds_are_rejected(size_px, mixed):
    with pytest.raises(ValueError) as failure:
        render_pptx([dict(elements=[sized_element(size_px, mixed)])])
    assert ("span#bounds" if mixed else "p#bounds") in str(failure.value)


@pytest.mark.parametrize("mixed", [False, True])
@pytest.mark.parametrize("size_px, expected_pt", [(2, 1), (2.5, 1.25), (8000, 4000)])
def test_representable_font_boundaries_and_fractional_sizes_are_retained(size_px, expected_pt, mixed):
    slide = reopen([dict(elements=[sized_element(size_px, mixed)])])
    text = text_shapes(slide)[0]
    assert text.text == "Boundary"
    assert text.text_frame.paragraphs[0].runs[0].font.size.pt == expected_pt
