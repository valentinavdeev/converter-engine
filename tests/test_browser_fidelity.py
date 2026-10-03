"""Synthetic HTML → Chromium → saved PPTX fidelity regressions."""
import math
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from html_to_pptx import convert


def html_file(tmp_path, content):
    source = tmp_path / 'input.html'
    source.write_text('''<!doctype html><meta charset="utf-8"><style>
*{box-sizing:border-box}body{margin:0}.slide{position:relative;width:1920px;height:1080px;
background:white;font-family:Arial;display:none}.slide.active{display:flex}
</style><section class="slide active">''' + content + '</section>', encoding='utf-8')
    return source


async def convert_html(tmp_path, content):
    source = html_file(tmp_path, content)
    target = tmp_path / 'output.pptx'
    await convert(str(source), str(target))
    return Presentation(target).slides[0]


@pytest.mark.asyncio
async def test_full_conversion_keeps_native_text_and_reading_surfaces(minimal_html: Path, tmp_pptx: Path):
    await convert(str(minimal_html), str(tmp_pptx))
    presentation = Presentation(tmp_pptx)
    expected = [
        ['Title Slide', 'Subtitle text here'],
        ['Content Slide', 'This is a minimal two-slide deck for testing the converter.'],
    ]
    assert [[s.text for s in slide.shapes if s.has_text_frame and s.text]
            for slide in presentation.slides] == expected
    assert [str(s.background.fill.fore_color.rgb) for s in presentation.slides] == ['1E293B', 'FFFFFF']
    title = next(s for s in presentation.slides[0].shapes if s.has_text_frame and s.text == 'Title Slide')
    assert title.text_frame.paragraphs[0].runs[0].font.size.pt == 32
    assert all(s.shape_type != MSO_SHAPE_TYPE.PICTURE for slide in presentation.slides for s in slide.shapes)


@pytest.mark.asyncio
async def test_preformatted_leaf_and_nested_boundaries_survive_browser(tmp_path):
    slide = await convert_html(tmp_path, '''
<span style="position:absolute;left:100px;top:100px;font-size:40px;white-space:pre">  A  B  </span>
<div style="position:absolute;left:100px;top:200px;font-size:40px;white-space:pre">  C <strong> D </strong> E  </div>
<span style="position:absolute;left:100px;top:300px;font-size:40px;white-space:normal">  Normal   text  </span>
''')
    texts = [s.text for s in slide.shapes if s.has_text_frame and s.text]
    assert texts == ['  A  B  ', '  C  D  E  ', 'Normal text']


@pytest.mark.asyncio
async def test_positive_subpixel_rules_survive_extraction_and_rendering(tmp_path):
    content = ''.join(
        f'<div style="position:absolute;left:100px;top:{100+i*60}px;width:200px;height:{h}px;background:rgb(255,0,0)"></div>'
        for i, h in enumerate([.5, 1, 2]))
    content += '<div style="position:absolute;left:400px;top:100px;width:.5px;height:200px;background:rgb(0,128,0)"></div>'
    content += '<div style="position:absolute;left:500px;top:100px;width:0;height:200px;background:rgb(0,0,255)"></div>'
    slide = await convert_html(tmp_path, content)
    geometry = sorted((round(s.width / 6350, 3), round(s.height / 6350, 3)) for s in slide.shapes)
    assert geometry == [(.5, 200), (200, .5), (200, 1), (200, 2)]
    assert all(s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE for s in slide.shapes)


@pytest.mark.asyncio
async def test_tracking_keeps_sign_and_inline_overrides(tmp_path):
    slide = await convert_html(tmp_path, '''
<span style="position:absolute;left:100px;top:100px;font-size:40px;letter-spacing:-2px;white-space:pre">TIGHT</span>
<div style="position:absolute;left:100px;top:200px;font-size:40px;letter-spacing:2px;white-space:pre">WIDE<span style="letter-spacing:-1px">TIGHT</span></div>
''')
    fields = {s.text: s for s in slide.shapes if s.has_text_frame and s.text}
    assert fields['TIGHT'].text_frame.paragraphs[0].runs[0]._r.rPr.get('spc') == '-100'
    runs = fields['WIDETIGHT'].text_frame.paragraphs[0].runs
    assert [(r.text, r._r.rPr.get('spc')) for r in runs] == [('WIDE', '100'), ('TIGHT', '-50')]
    assert all(r.font.size.pt == 20 for r in runs)


@pytest.mark.asyncio
async def test_off_center_rotation_keeps_tick_segments_connected_in_native_group(tmp_path):
    slide = await convert_html(tmp_path, '''
<div role="group" aria-label="Tick" style="position:absolute;left:0;top:0;width:500px;height:500px">
<div style="position:absolute;left:100px;top:100px;width:28.284271px;height:1.5px;background:green;transform:rotate(45deg);transform-origin:0 50%"></div>
<div style="position:absolute;left:120px;top:120px;width:42.426407px;height:1.5px;background:green;transform:rotate(-45deg);transform-origin:0 50%"></div>
</div>''')
    group = slide.shapes[0]
    assert group.shape_type == MSO_SHAPE_TYPE.GROUP
    assert group.name == 'Tick'

    def endpoint(shape, side):
        angle = math.radians(shape.rotation)
        x = (shape.left + shape.width / 2) / 6350
        y = (shape.top + shape.height / 2) / 6350
        half = shape.width / 6350 / 2
        return x + side * half * math.cos(angle), y + side * half * math.sin(angle)

    assert endpoint(group.shapes[0], -1) == pytest.approx((100, 100.75), abs=.03)
    assert endpoint(group.shapes[0], 1) == pytest.approx((120, 120.75), abs=.03)
    assert endpoint(group.shapes[1], -1) == pytest.approx((120, 120.75), abs=.03)
    assert endpoint(group.shapes[1], 1) == pytest.approx((150, 90.75), abs=.03)
